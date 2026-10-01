"""One fitted layout for PDF and DOCX; overflow never deletes resume facts."""

import io
import re
from dataclasses import dataclass, replace

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services.pdf_service import (
    _BULLET,
    _CONTACT,
    _HEADING,
    _LIST_SEPARATORS,
    _SECTIONS,
    THEMES,
    Theme,
    _plain,
    generate_resume_pdf,
)
from app.services.resume_structure import clean_placeholders


class PageOverflow(ValueError):
    def __init__(self, pages: int, target: int):
        self.pages = pages
        self.target = target
        super().__init__(
            f"This resume needs {pages} pages at readable font sizes. "
            "Prioritize older optional bullets or projects, or choose two pages. "
            "Contact, roles, dates, education and skills have not been removed."
        )


@dataclass(frozen=True)
class ResumeExport:
    markdown: str
    theme: Theme
    pdf: bytes
    page_count: int
    page_target: int


def fit_resume(
    markdown: str, full_name: str = "", template: str = "modern", page_target: int = 2
) -> ResumeExport:
    if page_target not in (1, 2):
        raise ValueError("page_target must be 1 or 2")
    if template not in THEMES:
        raise ValueError("Unknown resume template")
    text = clean_placeholders(markdown)
    original = THEMES[template]
    # Contact text is 0.6 pt below body in the existing renderer.
    original = replace(original, body_size=max(11, original.body_size))
    candidates = [original]
    candidates.extend(
        replace(original, section_gap=gap, leading=leading) for gap, leading in ((6, 1.2), (3, 1.1))
    )
    candidates.extend(replace(candidates[-1], body_size=size) for size in (10.8, 10.6))
    count = 0
    for theme in candidates:
        pdf = generate_resume_pdf(text, full_name, template, theme_override=theme, strict=True)
        with fitz.open(stream=pdf, filetype="pdf") as parsed:
            count = len(parsed)
        if count <= page_target:
            return ResumeExport(text, theme, pdf, count, page_target)
    raise PageOverflow(count, page_target)


def generate_resume_docx(layout: ResumeExport, full_name: str = "") -> bytes:
    document = Document()
    section = document.sections[0]
    theme = layout.theme
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    # ReportLab frames have 6pt padding inside the page margins.
    section.left_margin = section.right_margin = Inches(theme.margin_x) + Pt(6)
    section.top_margin = section.bottom_margin = Inches(theme.margin_y) + Pt(6)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman" if theme.regular.startswith("Times") else "Arial"
    normal.font.size = Pt(theme.body_size)
    normal.paragraph_format.line_spacing = Pt(theme.body_size * theme.leading)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.widow_control = False
    first = True
    in_header = True
    has_markdown_sections = any(_HEADING.match(raw.strip()) for raw in layout.markdown.splitlines())
    for raw in layout.markdown.splitlines():
        line = raw.strip()
        if not line or line in ("---", "***", "___"):
            continue
        heading = re.match(r"^(#{1,6})\s+", line)
        level = len(heading[1]) if heading else 0
        value = line[heading.end() :] if heading else line
        value = _plain(value)
        if first:
            paragraph = document.add_paragraph(full_name or value)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if theme.header_align == 1 else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.runs[0].bold = True
            paragraph.runs[0].font.size = Pt(theme.name_size)
            paragraph.paragraph_format.line_spacing = Pt(theme.name_size * 1.15)
            paragraph.paragraph_format.space_after = Pt(3)
            first = False
            if level == 1 or (full_name or value).casefold() == value.casefold():
                continue
        title = value.rstrip(":")
        caps_heading = (
            not has_markdown_sections
            and title.isupper()
            and 2 < len(title) < 40
            and not _LIST_SEPARATORS.search(title)
        )
        if title.casefold() in _SECTIONS or (level == 2 and len(title) < 50) or caps_heading:
            in_header = False
            paragraph = document.add_paragraph(title.upper())
            paragraph.paragraph_format.space_before = Pt(theme.section_gap)
            paragraph.paragraph_format.space_after = Pt(6 + theme.heading_rule)
            paragraph.paragraph_format.line_spacing = Pt(theme.heading_size * 1.25)
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.runs[0]
            run.bold = True
            run.font.size = Pt(theme.heading_size)
            run.font.color.rgb = RGBColor.from_string(theme.ink.lstrip("#"))
        else:
            bullet = _BULLET.match(value)
            paragraph = document.add_paragraph("• " + value[bullet.end() :] if bullet else value)
            if in_header and _CONTACT.search(value):
                paragraph.runs[0].font.size = Pt(theme.body_size - 0.6)
                paragraph.paragraph_format.line_spacing = Pt(theme.body_size * 1.35)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if theme.header_align == 1
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
            elif in_header and level == 0 and not bullet and len(value) < 90:
                paragraph.runs[0].font.size = Pt(theme.body_size + 1)
                paragraph.paragraph_format.line_spacing = Pt((theme.body_size + 1) * 1.3)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if theme.header_align == 1
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
            else:
                in_header = False
                if bullet:
                    paragraph.paragraph_format.left_indent = Pt(12)
                    paragraph.paragraph_format.first_line_indent = Pt(-10)
                    paragraph.paragraph_format.space_after = Pt(1.2)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
